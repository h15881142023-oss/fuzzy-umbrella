"""用本机「新商考核预警数据」Excel 更新新商评当期（不核环比）。"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

import sync_xinshang_from_chuxin as xin  # noqa: E402

CACHE = ROOT / "data" / "xinshang"
CITIES = xin.CITIES
HTMLS = xin.HTMLS
EXACT_REL = (
    Path("我的文档")
    / "xwechat_files"
    / "wxid_acpkrxuu0ted22_f987"
    / "temp"
    / "RWTemp"
    / "2026-09"
    / "f7219ce6041b6c82ef27149b5631a659"
    / "新商考核预警数据_20260908_113823.xlsx"
)
EXACT_NAME = "新商考核预警数据_20260908_113823.xlsx"
FILE_DATE_RE = re.compile(r"(\d{8})")
ISO_RE = re.compile(r"(20\d{2}-\d{2}-\d{2})")

FIELD_ALIASES = {
    "城市名": "城市",
    "城市级别": "城市等级",
    "综合排名": "大盘预警",
    "加权排名": "加权排名区间",
    "餐饮渗透率": "餐饮商家渗透率指标值-外卖",
    "餐饮商家渗透率": "餐饮商家渗透率指标值-外卖",
    "餐饮渗透率排名": "餐饮商家渗透率-外卖",
    "餐饮订单量完成率": "市场开发率（订单）指标值-外卖",
    "餐饮订单量完成率排名": "市场开发率（订单）-外卖",
    "餐饮交易额完成率": "市场开发率（实付）指标值-外卖",
    "餐饮交易额完成率排名": "市场开发率（实付）-外卖",
    "市场开发率差值": "市场开发率（订单）指标值-外卖",
    "市场开发率": "市场开发率（订单）指标值-外卖",
    "市场开发率_GMV差值": "市场开发率（实付）指标值-外卖",
    "团购市场开发率变动": "市场开发率指标值-团购",
    "团购市场开发率排名": "市场开发率-团购",
    "优质商家渗透率": "优质商家渗透率指标值-团购",
    "渗透率排名": "优质商家渗透率-团购",
    "非餐YOY": "YoY指标值-零售",
    "零售_YoY": "YoY指标值-零售",
    "日均零售YOY": "YoY指标值-零售",
    "零售_YoY排名": "YoY-零售预警",
    "优质仓达标情况": "优质仓数达标情况",
    "外卖模块预警": "外卖能力预警",
    "团购模块预警": "团购能力预警",
    "履约模块预警": "履约能力预警",
    "零售模块预警": "零售能力预警",
    "组织模块预警": "组织能力预警",
    "商业增值模块预警": "商业增值能力预警",
    "用户体验模块预警": "用户体验能力预警",
    "综合治理模块预警": "综合治理能力预警",
    "推单完成率": "推单完成率指标值-履约",
    "推单完成率_调度后": "推单完成率指标值-履约",
    "推单排名": "推单完成率排名-履约",
    "压力天出勤率": "压力天出勤率",
    "超45分钟订单占比": "超45分钟订单占比指标值-履约",
    "超45分钟订单排名": "超45分钟订单占比-履约",
    "外卖货币化率": "外卖货币化率指标值-商业增值",
    "团购货币化率": "团购货币化率指标值-商业增值",
    "商业增值_外卖货币化率排名": "商业增值_外卖货币化率排名",
    "商业增值_团购货币化率排名": "商业增值_团购货币化率排名",
    "商家投诉差值": "用户体验_用户投诉商家问题万服差值",
    "商家投诉排名": "用户体验_用户投诉商家问题万服差值排名",
    "履约投诉差值": "用户体验_用户投诉履约问题万服差值",
    "履约投诉排名": "用户体验_用户投诉履约问题万服差值排名",
}


def _exists(p: Path) -> bool:
    try:
        return p.is_file() and p.stat().st_size > 100
    except OSError:
        return False


def candidate_paths(explicit: str | None = None) -> list[Path]:
    import os

    out: list[Path] = [
        Path(r"d:\我的文档\xwechat_files\wxid_acpkrxuu0ted22_f987\temp\RWTemp\2026-09\f7219ce6041b6c82ef27149b5631a659")
        / EXACT_NAME,
        Path(r"D:\我的文档") / EXACT_REL,
        Path(r"C:\我的文档") / EXACT_REL,
    ]
    if explicit:
        out.insert(0, Path(explicit))
    env = os.environ.get("XINSHANG_EXCEL")
    if env:
        out.insert(0, Path(env))
    out.extend(
        [
            ROOT / "data" / "xinshang" / EXACT_NAME,
            ROOT / EXACT_NAME,
            Path("/home/ubuntu/.cursor/projects/workspace/uploads") / "_________20260908_113823_7c4a.xlsx",
            Path.home() / "Documents" / EXACT_NAME,
            Path.home() / "Desktop" / EXACT_NAME,
            Path.home() / "Downloads" / EXACT_NAME,
            Path.home() / "我的文档" / EXACT_NAME,
        ]
    )
    seen: list[Path] = []
    for p in out:
        if p not in seen:
            seen.append(p)
    return seen


def download_excel_from_repo(dest: Path) -> Path | None:
    import urllib.request

    rel = f"data/xinshang/{EXACT_NAME}"
    urls = [
        f"https://ghproxy.net/https://raw.githubusercontent.com/h15881142023-oss/fuzzy-umbrella/cursor/cz1-merchant-dashboard-74a9/{rel}",
        f"https://raw.githubusercontent.com/h15881142023-oss/fuzzy-umbrella/cursor/cz1-merchant-dashboard-74a9/{rel}",
    ]
    dest.parent.mkdir(parents=True, exist_ok=True)
    for url in urls:
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "cz1-xinshang"})
            with urllib.request.urlopen(req, timeout=60) as resp:
                data = resp.read()
            if data and len(data) > 200 and data[:2] == b"PK":
                dest.write_bytes(data)
                return dest
        except Exception:
            continue
    return None


def find_excel(explicit: str | None = None) -> Path | None:
    for p in candidate_paths(explicit):
        if _exists(p):
            return p
    # 限定目录里找最新「新商考核预警数据*.xlsx」
    search_roots = [
        Path(r"d:\我的文档\xwechat_files"),
        Path(r"D:\我的文档\xwechat_files"),
        Path.home() / "Documents" / "xwechat_files",
        Path.home() / "我的文档" / "xwechat_files",
        ROOT / "data" / "xinshang",
        Path("/home/ubuntu/.cursor/projects/workspace/uploads"),
        Path.home() / "Desktop",
        Path.home() / "Downloads",
        Path.home() / "Documents",
    ]
    found: list[Path] = []
    for root in search_roots:
        if not root.exists():
            continue
        try:
            found.extend(root.rglob("新商考核预警数据*.xlsx"))
        except OSError:
            continue
    found = [p for p in found if _exists(p) and not p.name.startswith("~$")]
    if not found:
        return None
    found.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    if found:
        return found[0]
    return download_excel_from_repo(CACHE / EXACT_NAME)


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    return "" if s.lower() in {"none", "nan", "nat", "null"} else s


def normalize_header(v) -> str:
    return re.sub(r"\s+", "", cell_str(v))


def looks_like_header(vals: list[str]) -> bool:
    joined = "".join(vals)
    return "城市" in joined and any(k in joined for k in ("预警", "区域", "等级", "渗透", "日期"))


def detect_header_row(rows: list[list]) -> int:
    best = 0
    for i, row in enumerate(rows[:12]):
        vals = [cell_str(c) for c in row]
        if looks_like_header(vals):
            return i
        if "城市" in "".join(vals) and best == 0:
            best = i
    return best


def sheet_records(ws) -> tuple[list[str], list[dict]]:
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    if not rows:
        return [], []
    hi = detect_header_row(rows)
    cols = []
    seen = {}
    for raw in rows[hi]:
        name = normalize_header(raw) or f"col{len(cols)}"
        n = seen.get(name, 0) + 1
        seen[name] = n
        cols.append(name if n == 1 else f"{name}_{n}")
    recs = []
    for raw in rows[hi + 1 :]:
        if not any(cell_str(x) for x in raw):
            continue
        recs.append({cols[i]: raw[i] if i < len(raw) else None for i in range(len(cols))})
    return cols, recs


def canon_city(v) -> str | None:
    s = xin.clean_city_name(v)
    if not s:
        return None
    for name, keys in (
        ("彭州市", ("彭州市", "彭州")),
        ("仁寿县", ("仁寿县", "仁寿")),
        ("合江县", ("合江县", "合江")),
        ("南溪", ("南溪区", "南溪县", "南溪")),
        ("叙永", ("叙永县", "叙永")),
    ):
        if any(k in s for k in keys):
            return name
    return None


def pick_date(values: list[str], filename: str) -> str:
    days = []
    for v in values:
        m = ISO_RE.search(cell_str(v))
        if m:
            days.append(m.group(1))
        s = cell_str(v)
        if len(s) >= 10 and s[4] == "-" and s[0].isdigit():
            days.append(s[:10])
    days = sorted({d for d in days if d >= "2026-01-01"})
    if days:
        return days[-1]
    m = FILE_DATE_RE.search(filename)
    if m:
        raw = m.group(1)
        return f"{raw[:4]}-{raw[4:6]}-{raw[6:8]}"
    return datetime.now().strftime("%Y-%m-%d")


def alias_row(row: dict) -> dict:
    out = dict(row)
    for src, dst in FIELD_ALIASES.items():
        if src in row and (dst not in out or xin.blank(out.get(dst))):
            out[dst] = row.get(src)
    city = canon_city(row.get("城市") or row.get("城市名") or row.get("城市名称") or row.get("city"))
    if city:
        out["城市"] = city
    return out


def merge_city_maps(maps: list[dict[str, dict]]) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for mp in maps:
        for city, row in mp.items():
            cur = out.setdefault(city, {"城市": city})
            for k, v in row.items():
                if not xin.blank(v):
                    cur[k] = v
    return out


def inspect_book(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = []
    city_maps = []
    date_vals = []
    for name in wb.sheetnames:
        ws = wb[name]
        cols, recs = sheet_records(ws)
        mapped = {}
        for rec in recs:
            row = alias_row(rec)
            city = canon_city(rec.get("城市") or rec.get("城市名") or row.get("城市"))
            if not city:
                raw = xin.clean_city_name(rec.get("城市") or rec.get("城市名") or "")
                city = raw or None
            if not city:
                continue
            row["城市"] = city
            mapped[city] = row
            for key in ("日期", "最新数据日期", "数据日期", "本期日期"):
                if rec.get(key) is not None:
                    date_vals.append(cell_str(rec.get(key)))
        city_maps.append(mapped)
        sample = {}
        for city in CITIES:
            if city in mapped:
                sample[city] = {k: cell_str(v)[:40] for k, v in list(mapped[city].items())[:18]}
        sheets.append(
            {
                "name": name,
                "cols": cols[:40],
                "n": len(recs),
                "hitCities": sorted(mapped),
                "sample": sample,
            }
        )
    wb.close()
    merged = merge_city_maps(city_maps)
    return {
        "file": str(path),
        "size": path.stat().st_size,
        "sheets": sheets,
        "mergedCities": sorted(merged),
        "periodGuess": pick_date(date_vals + [path.name], path.name),
        "merged": merged,
    }


def fetch_prev_from_metabase(day: str) -> tuple[str | None, dict, dict, dict, dict, dict, dict]:
    """上期汇总仍走考核日；当期外卖/团购模块补交易商家数、公海、团购子项。彭州走川藏二区。"""
    prev_day = xin.prev_assessment_date(day)
    prev: dict = {}
    board: dict = {}
    waimai_prev: dict = {}
    tuango_prev: dict = {}
    if prev_day:
        tables = {}
        for name, spec in xin.CARDS.items():
            cols, rows = xin.query_card_regions(spec, prev_day)
            tables[name] = {"cols": cols, "rows": rows}
        prev = xin.rows_to_city_map(tables["summary"]["cols"], tables["summary"]["rows"])
        board = xin.pick_latest_board(tables["cityboard"]["cols"], tables["cityboard"]["rows"])
        waimai_prev = xin.rows_to_city_map(tables["waimai"]["cols"], tables["waimai"]["rows"])
        tuango_prev = xin.rows_to_city_map(tables["tuango"]["cols"], tables["tuango"]["rows"])
    waimai_cur = xin.rows_to_city_map(*xin.query_card_regions(xin.CARDS["waimai"], day))
    tuango_cur = xin.rows_to_city_map(*xin.query_card_regions(xin.CARDS["tuango"], day))
    return prev_day, prev, board, waimai_prev, tuango_prev, waimai_cur, tuango_cur


def update_peer_compare(data: dict, day: str, prev_day: str | None, excel_all: dict[str, dict], prev: dict) -> None:
    import sync_peer_compare_from_chuxin as peer

    period, prev_d, summary, summary_prev, modules, dump = peer.fetch_all(day)
    if prev_day and not prev_d:
        prev_d = prev_day
    if not summary_prev and prev:
        summary_prev = prev
    for city, row in excel_all.items():
        cur = dict(summary.get(city) or {})
        cur.update(row)
        summary[city] = cur
    payload = peer.build_payload(period, prev_d, summary, summary_prev, modules, dump)
    data["peerCompare"] = payload


def apply_excel(day: str, merged: dict[str, dict], inspect: dict) -> dict:
    missing = [c for c in CITIES if c not in merged]
    if missing:
        raise RuntimeError(f"Excel 缺城: {missing}; got {[c for c in CITIES if c in merged]}")

    prev_day, prev, board, waimai_prev, tuango_prev, waimai_cur, tuango_cur = fetch_prev_from_metabase(day)

    html = HTMLS[0].read_text(encoding="utf-8")
    start, end, data = xin.extract_data_json(html)
    scraped_at = datetime.now(timezone.utc).isoformat()
    data.setdefault("meta", {})
    data["meta"]["period"] = f"{day[:4]}年{int(day[5:7])}月"
    data["meta"]["dataDate"] = day
    data["meta"]["prevPeriod"] = prev_day
    data["meta"]["scrapedAt"] = scraped_at
    data["meta"]["testSyncAt"] = scraped_at
    data["meta"]["obsTitle"] = "新商考核预警数据 Excel"
    data["meta"]["obsUpdatedAt"] = day
    data["meta"]["excelFile"] = Path(inspect["file"]).name
    data["source"] = {
        **(data.get("source") or {}),
        "dashboard": "新商考核预警数据",
        "tab": Path(inspect["file"]).name,
        "periodDate": day,
        "prevDate": prev_day,
        "fromExcel": True,
        "prevSource": "metabase",
    }

    online_map = xin.load_powerbi_online()
    by_name = {c["name"]: c for c in (data.get("cities") or [])}
    new_cities = []
    for name in CITIES:
        city = by_name.get(name) or {"name": name}
        city["name"] = name
        row = merged[name]
        xin.apply_city(
            city,
            row,
            prev.get(name),
            board.get(name) or row,
            waimai_cur.get(name) or {},
            tuango_cur.get(name) or {},
            online_map,
            waimai_prev.get(name),
            tuango_prev.get(name),
        )
        city["dataDate"] = day
        new_cities.append(city)
    data["cities"] = new_cities
    xin.rename_waimai_layouts(data.get("layouts") or {})
    xin.enable_layout_mom(data.get("layouts") or {})
    update_peer_compare(data, day, prev_day, merged, prev)

    new_json = json.dumps(data, ensure_ascii=False, indent=2)
    new_html = html[:start] + new_json + html[end:]
    footer_new = (
        f"能力指标同步自「新商考核预警数据」（{day}，上期 {prev_day or '—'} / Metabase）；"
        "测评成绩另见集合「新商评测试结果」"
    )
    new_html = re.sub(
        r"能力指标同步自[^<]+；测评成绩另见集合「新商评测试结果」",
        footer_new,
        new_html,
    )
    for p in HTMLS:
        p.write_text(new_html, encoding="utf-8")
    return {
        "ok": True,
        "usedExcel": True,
        "date": day,
        "prev": prev_day,
        "file": Path(inspect["file"]).name,
        "cities": {
            n: {
                "level": merged[n].get("城市等级"),
                "market": merged[n].get("大盘预警"),
                "waimai": merged[n].get("外卖能力预警"),
            }
            for n in CITIES
        },
        "sheets": [s["name"] for s in inspect.get("sheets") or []],
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", help="Excel 路径")
    ap.add_argument("--inspect-only", action="store_true")
    args = ap.parse_args()

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print(json.dumps({"ok": False, "usedExcel": False, "error": "缺少 openpyxl"}, ensure_ascii=False))
        return 2

    path = find_excel(args.xlsx)
    if not path:
        print(
            json.dumps(
                {"ok": False, "usedExcel": False, "error": "未找到新商考核预警数据 Excel", "tried": [str(p) for p in candidate_paths(args.xlsx)[:8]]},
                ensure_ascii=False,
            )
        )
        return 3

    CACHE.mkdir(parents=True, exist_ok=True)
    dest = CACHE / path.name
    try:
        dest.write_bytes(path.read_bytes())
    except OSError:
        dest = path

    inspect = inspect_book(path)
    slim = {k: inspect[k] for k in ("file", "size", "sheets", "mergedCities", "periodGuess")}
    (CACHE / "excel_inspect.json").write_text(json.dumps(slim, ensure_ascii=False, indent=2), encoding="utf-8")

    if args.inspect_only:
        print(json.dumps({"ok": True, "usedExcel": False, **slim}, ensure_ascii=False))
        return 0

    try:
        result = apply_excel(inspect["periodGuess"], inspect["merged"], inspect)
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "usedExcel": False, "error": str(exc), **slim}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
